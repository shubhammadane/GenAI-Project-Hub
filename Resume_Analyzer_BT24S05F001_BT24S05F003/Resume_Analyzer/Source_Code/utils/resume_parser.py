import io
import re
from typing import Optional

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pypdf, with fallback to pypdfium2."""
    text = ""
    # Method 1: pypdf
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return clean_extracted_text(text)
    except Exception as e:
        print(f"pypdf extraction error: {e}")

    # Method 2: pypdfium2 (handles diverse PDF specs)
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(file_bytes)
        for page in pdf:
            textpage = page.get_textpage()
            page_text = textpage.get_text_range()
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return clean_extracted_text(text)
    except Exception as e:
        print(f"pypdfium2 extraction error: {e}")

    return clean_extracted_text(text)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from docx file."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        # Also check tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text += "\n" + row_text
        return clean_extracted_text(text)
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"

def extract_text(uploaded_file) -> str:
    """Universal extractor based on file extension or type."""
    if uploaded_file is None:
        return ""
    
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx") or filename.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    elif filename.endswith(".txt"):
        try:
            return clean_extracted_text(file_bytes.decode("utf-8", errors="ignore"))
        except Exception:
            return clean_extracted_text(str(file_bytes))
    else:
        # Try UTF-8 decode
        try:
            return clean_extracted_text(file_bytes.decode("utf-8", errors="ignore"))
        except Exception:
            return ""

def clean_extracted_text(text: str) -> str:
    """Clean whitespace and formatting artifacts."""
    if not text:
        return ""
    # Normalize excessive newlines and spaces
    text = re.sub(r"\r\n", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
